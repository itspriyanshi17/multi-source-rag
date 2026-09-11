"""One-off converter: REPORT.md -> REPORT.docx. Not part of the app; run manually when the
report changes. Handles the specific markdown subset used in REPORT.md (headings, pipe tables,
bullet lists, bold spans, horizontal rules, fenced code blocks).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REPORT_MD = PROJECT_ROOT / "REPORT.md"
REPORT_DOCX = PROJECT_ROOT / "REPORT.docx"

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_bold_aware_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        run = p.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])
    return p


def parse_table(lines, start_idx):
    header = [c.strip() for c in lines[start_idx].strip().strip("|").split("|")]
    rows = []
    i = start_idx + 2  # skip header and separator row
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    return header, rows, i


def _starts_new_block(stripped: str) -> bool:
    return (
        not stripped
        or stripped.startswith("#")
        or stripped.startswith("|")
        or stripped.startswith("```")
        or stripped == "---"
        or stripped.startswith("- ")
        or stripped.startswith("* ")
        or bool(re.match(r"^\d+\.\s", stripped))
    )


def _collect_wrapped_text(lines, start_idx, first_line_text):
    """Joins `first_line_text` with any following lines that are markdown "soft wraps" of the
    same paragraph/list item (non-blank, not starting a new block) into one logical string."""
    parts = [first_line_text]
    i = start_idx + 1
    while i < len(lines) and not _starts_new_block(lines[i].strip()):
        parts.append(lines[i].strip())
        i += 1
    return " ".join(parts), i


def build_docx():
    text = REPORT_MD.read_text(encoding="utf-8")
    lines = text.split("\n")

    doc = Document()
    for style_name, size in [("Normal", 11), ("Heading 1", 20), ("Heading 2", 16), ("Heading 3", 13)]:
        style = doc.styles[style_name]
        style.font.size = Pt(size)

    in_code_block = False
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            in_code_block = not in_code_block
            i += 1
            continue

        if in_code_block:
            p = doc.add_paragraph(line)
            p.style = doc.styles["Normal"]
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            i += 1
            continue

        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("_" * 60).alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading_text = stripped.lstrip("#").strip()
            level = min(max(level, 1), 3)
            doc.add_heading(heading_text, level=level)
            i += 1
            continue

        if stripped.startswith("|"):
            header, rows, next_i = parse_table(lines, i)
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                table.rows[0].cells[j].text = h
                for p in table.rows[0].cells[j].paragraphs:
                    for r in p.runs:
                        r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for j, val in enumerate(row):
                    if j < len(cells):
                        cells[j].text = val
            i = next_i
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            joined, next_i = _collect_wrapped_text(lines, i, stripped[2:])
            add_bold_aware_paragraph(doc, joined, style="List Bullet")
            i = next_i
            continue

        if re.match(r"^\d+\.\s", stripped):
            content = re.sub(r"^\d+\.\s", "", stripped)
            joined, next_i = _collect_wrapped_text(lines, i, content)
            add_bold_aware_paragraph(doc, joined, style="List Number")
            i = next_i
            continue

        joined, next_i = _collect_wrapped_text(lines, i, stripped)
        add_bold_aware_paragraph(doc, joined)
        i = next_i

    doc.save(REPORT_DOCX)
    print(f"Wrote {REPORT_DOCX}")


if __name__ == "__main__":
    build_docx()
